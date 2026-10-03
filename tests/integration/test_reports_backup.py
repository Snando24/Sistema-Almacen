"""Reportes, exportaciones, configuración, auditoría y respaldos."""

from __future__ import annotations

import sqlite3
import zipfile
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest

from sisalmacen.application.reports import REPORTS, ReportParams
from sisalmacen.domain.errors import PermissionDenied, ValidationError
from sisalmacen.domain.inventory import (
    AuditFilter,
    MovementLine,
    MovementRequest,
    ProductFilter,
)

D = Decimal


def _seed(services, make_product):  # type: ignore[no-untyped-def]
    one = make_product("P-1", precio_compra=D("2.50"), stock_minimo=D(5))
    two = make_product("P-2", precio_compra=D("4"))
    for product_id, qty in ((one, "4"), (two, "10")):
        services.movements.register(
            MovementRequest("ENT_COMPRA", date(2026, 1, 10), (MovementLine(product_id, D(qty)),))
        )
    return one, two


def test_all_reports_build_with_known_data(services, make_product) -> None:  # type: ignore[no-untyped-def]
    one, _two = _seed(services, make_product)
    params = ReportParams(
        producto_id=one, fecha_desde=date(2026, 1, 1), fecha_hasta=date(2026, 1, 31)
    )
    built = {key: services.reports.build(key, params) for key, _title in REPORTS}

    assert len(built["inventario_general"].rows) == 2
    assert [r[0] for r in built["stock_bajo"].rows] == ["P-1"]
    assert built["sin_stock"].rows == []
    assert built["por_categoria"].rows == [["General", 2, 2, D("50.0")]]
    assert len(built["movimientos_periodo"].rows) == 2
    assert len(built["historial_producto"].rows) == 1
    valued = built["valorizado"]
    assert valued.totals is not None
    assert valued.totals[-1] == D("50.0")  # 4×2.50 + 10×4


def test_product_history_requires_a_product(services) -> None:  # type: ignore[no-untyped-def]
    with pytest.raises(ValidationError):
        services.reports.build("historial_producto", ReportParams())


def test_exports_create_files_and_are_audited(services, make_product, tmp_path: Path) -> None:  # type: ignore[no-untyped-def]
    _seed(services, make_product)
    data = services.reports.products_listing(ProductFilter())

    csv_path = tmp_path / "productos.csv"
    xlsx_path = tmp_path / "productos.xlsx"
    pdf_path = tmp_path / "productos.pdf"
    for fmt, path in (("csv", csv_path), ("xlsx", xlsx_path), ("pdf", pdf_path)):
        services.reports.export(data, fmt, path)
        assert path.stat().st_size > 0

    lines = csv_path.read_text(encoding="utf-8-sig").splitlines()
    assert lines[0].startswith("Código,Nombre")
    assert len(lines) == 3

    from openpyxl import load_workbook

    sheet = load_workbook(xlsx_path).active
    header_row = next(r for r in sheet.iter_rows() if r[0].value == "Código")
    assert sheet.auto_filter.ref is not None
    assert header_row[0].row < sheet.max_row
    assert pdf_path.read_bytes().startswith(b"%PDF")

    audited = services.audit.search(AuditFilter(accion="EXPORTAR")).items
    assert len(audited) == 3


def test_settings_update_is_audited_and_validated(services) -> None:  # type: ignore[no-untyped-def]
    services.settings.update({"empresa.nombre": "R&R Prueba", "inventario.usar_precios": "0"})
    assert services.settings.read_all()["empresa.nombre"] == "R&R Prueba"
    entry = services.audit.search(AuditFilter(accion="CONFIGURAR")).items[0]
    assert "R&R Prueba" in (entry.valor_nuevo or "")
    with pytest.raises(ValidationError):
        services.settings.update({"importacion.max_filas": "abc"})
    with pytest.raises(ValidationError):
        services.settings.update({"clave.inventada": "1"})
    services.settings.update({"ui.tema": "oscuro"})
    assert services.settings.read_all()["ui.tema"] == "oscuro"
    with pytest.raises(ValidationError):
        services.settings.update({"ui.tema": "sistema"})


def test_prices_hidden_when_disabled(services, make_product) -> None:  # type: ignore[no-untyped-def]
    product_id = make_product("P-1", precio_compra=D("3"))
    services.settings.update({"inventario.usar_precios": "0"})
    assert services.products.get(product_id).precio_compra is None
    with pytest.raises(PermissionDenied):
        services.reports.build("valorizado", ReportParams())


def test_audit_is_append_only(app_context, services, make_product) -> None:  # type: ignore[no-untyped-def]
    make_product("P-1")
    connection = sqlite3.connect(app_context.paths.database_path)
    try:
        with pytest.raises(sqlite3.DatabaseError):
            connection.execute("UPDATE auditoria SET accion = 'X'")
        with pytest.raises(sqlite3.DatabaseError):
            connection.execute("DELETE FROM auditoria")
    finally:
        connection.close()


def test_backup_create_and_restore_roundtrip(
    services, make_product, app_context, tmp_path: Path
) -> None:  # type: ignore[no-untyped-def]
    make_product("P-1")
    archive, sha256 = services.backup.create(tmp_path / "bk")
    with zipfile.ZipFile(archive) as bundle:
        assert {"sisalmacen.sqlite3", "manifest.json"} <= set(bundle.namelist())
    assert len(sha256) == 64
    assert services.backup.reminder() is None
    assert services.audit.search(AuditFilter(accion="RESPALDAR")).total == 1

    make_product("P-2")
    assert services.products.search(ProductFilter()).total == 2

    previous = services.backup.restore(archive)

    assert previous.exists()
    assert services.products.search(ProductFilter()).total == 1
    assert services.movements.verify_consistency() == []


def test_restore_rejects_corrupt_or_foreign_archives(services, tmp_path: Path) -> None:  # type: ignore[no-untyped-def]
    not_zip = tmp_path / "x.zip"
    not_zip.write_text("no es un zip")
    with pytest.raises(ValidationError) as error:
        services.backup.restore(not_zip)
    assert error.value.code == "BACKUP_INVALIDO"

    archive, _ = services.backup.create(tmp_path / "bk")
    tampered = tmp_path / "tampered.zip"
    with zipfile.ZipFile(archive) as source, zipfile.ZipFile(tampered, "w") as target:
        for name in source.namelist():
            content = source.read(name)
            target.writestr(name, content + b"x" if name.endswith(".sqlite3") else content)
    with pytest.raises(ValidationError) as hash_error:
        services.backup.restore(tampered)
    assert hash_error.value.code == "BACKUP_INVALIDO"


def test_backup_reminder_when_never_done(services) -> None:  # type: ignore[no-untyped-def]
    assert services.backup.reminder() is not None
