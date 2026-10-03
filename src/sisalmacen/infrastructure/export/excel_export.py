"""Exportación a Excel (RF-091)."""

from __future__ import annotations

from collections.abc import Mapping
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from sisalmacen.domain.inventory import ReportData

_HEADER_COLOR = "D1121E"
_NUMBER_FORMATS = {"qty": "#,##0.###", "money": "#,##0.00##", "int": "#,##0"}


class ExcelReportWriter:
    """Implementa `ReportWriterPort` para XLSX con encabezado de empresa y autofiltro."""

    def write(self, path: Path, data: ReportData, *, company: Mapping[str, str]) -> None:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None  # noqa: S101
        sheet.title = data.title[:31]

        company_name = company.get("empresa.nombre", "")
        sheet.append([company_name])
        sheet.append([data.title])
        if data.filters:
            sheet.append(["Filtros: " + " | ".join(data.filters)])
        sheet.append([f"Generado: {datetime.now().astimezone():%Y-%m-%d %H:%M}"])
        sheet["A1"].font = Font(bold=True, size=14)
        sheet["A2"].font = Font(bold=True, size=12)

        header_row = sheet.max_row + 1
        sheet.append(data.headers)
        fill = PatternFill("solid", fgColor=_HEADER_COLOR)
        for column in range(1, len(data.headers) + 1):
            cell = sheet.cell(row=header_row, column=column)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in data.rows:
            sheet.append(_convert_row(row, data.column_types))
            self._format_row(sheet, sheet.max_row, data.column_types)
        last_data_row = sheet.max_row
        if data.totals is not None:
            sheet.append(_convert_row(data.totals, data.column_types))
            self._format_row(sheet, sheet.max_row, data.column_types)
            for column in range(1, len(data.headers) + 1):
                sheet.cell(row=sheet.max_row, column=column).font = Font(bold=True)

        sheet.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(len(data.headers))}{max(last_data_row, header_row)}"
        )
        sheet.freeze_panes = f"A{header_row + 1}"
        for index, header in enumerate(data.headers, start=1):
            longest = max(
                [len(str(header))]
                + [len(str(row[index - 1])) for row in data.rows if index - 1 < len(row)]
            )
            sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 10), 50)
        workbook.save(path)

    @staticmethod
    def _format_row(sheet: Worksheet, row_number: int, types: list[str]) -> None:
        for column, kind in enumerate(types, start=1):
            number_format = _NUMBER_FORMATS.get(kind)
            if number_format:
                sheet.cell(row=row_number, column=column).number_format = number_format


def _convert_row(row: list[object], types: list[str]) -> list[object]:
    converted: list[object] = []
    for value, kind in zip(row, types, strict=False):
        if kind == "date" and isinstance(value, str) and value:
            try:
                converted.append(date.fromisoformat(value))
                continue
            except ValueError:
                pass
        converted.append(value)
    return converted
